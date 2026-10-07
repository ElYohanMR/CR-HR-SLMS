"""Portable inventory files. One row per physical locker; no employee data."""

import csv
import math
from pathlib import Path

HEADERS = [
    "id_bloque",
    "version_bloque",
    "plano",
    "edificio",
    "bloque",
    "lado",
    "ubicacion",
    "filas",
    "columnas",
    "x",
    "y",
    "rotacion",
    "fila",
    "columna",
    "numero",
]
MAX_ROWS = 20000


def text(value):
    return "" if value is None else str(value).strip()


def export_rows(snapshot):
    buildings = {b["id"]: b["code"] for b in snapshot["buildings"]}
    blocks = {b["id"]: b for b in snapshot["blocks"]}
    result = []
    for locker in sorted(snapshot["lockers"], key=lambda item: (item["block_id"], item["row"], item["column"])):
        b = blocks[locker["block_id"]]
        values = [
            b["id"],
            b["version"],
            b["plan_id"],
            buildings[b["building_id"]],
            b["label"],
            b["side"],
            b["location"],
            b["rows"],
            b["columns"],
            b["x"],
            b["y"],
            b["rotation"],
            locker["row"] + 1,
            locker["column"] + 1,
            locker["number"],
        ]
        result.append(dict(zip(HEADERS, map(text, values))))
    return result


def write_inventory(path, rows):
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        book = Workbook()
        sheet = book.active
        sheet.title = "Inventario"
        sheet.append(HEADERS)
        for row in rows:
            sheet.append([text(row.get(key)) for key in HEADERS])
        for cells in sheet:
            for cell in cells:
                cell.data_type = "s"
                cell.number_format = "@"
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="444444")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, key in enumerate(HEADERS, 1):
            sheet.column_dimensions[get_column_letter(index)].width = 39 if key == "id_bloque" else 18
        guide = book.create_sheet("Instrucciones")
        for line in [
            "Una fila por casillero; no se importan empleados ni asignaciones.",
            "Para crear bloques: id_bloque y version_bloque vacíos. El plano debe existir (ID o nombre único).",
            "Filas/columnas: tamaño del bloque. Fila/columna: posición del casillero, empezando en 1.",
            "X/Y: coordenadas de 0 a 1000 relativas al ancho/alto del plano. Rotación: -360 a 360.",
            "Ubicación EXT o INT. Lado N, S, E u O. Número físico único, conservado como texto.",
            "Para actualizar: conserva ID y versión exportados; modifica solo bloque (nombre), X, Y o rotación.",
            "Repite los mismos datos de bloque en todas sus filas. No cambies su numeración ni dimensiones.",
            "Los casilleros nuevos quedan pendientes de verificación física. Revisa Solicitudes al sincronizar.",
            "La importación es completa: si hay un conflicto no se aplica ningún cambio del archivo.",
        ]:
            guide.append([line])
        guide.column_dimensions["A"].width = 130
        book.save(path)
        book.close()
    elif path.suffix.lower() == ".csv":

        def safe(value):
            s = text(value)
            return "'" + s if s.startswith(("=", "+", "-", "@", "\t", "\r")) else s

        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, HEADERS)
            writer.writeheader()
            writer.writerows({key: safe(row.get(key)) for key in HEADERS} for row in rows)
    else:
        raise ValueError("Selecciona un archivo .xlsx o .csv.")


def read_inventory(path):
    path = Path(path)
    if path.stat().st_size > 20_000_000:
        raise ValueError("El archivo supera 20 MB.")
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            sheet = book["Inventario"] if "Inventario" in book.sheetnames else book.worksheets[0]
            rows = []
            for index, cells in enumerate(sheet.iter_rows()):
                if index > MAX_ROWS:
                    raise ValueError(f"Máximo {MAX_ROWS} casilleros por archivo.")
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError(f"Fila {index + 1}: utiliza valores, no fórmulas.")
                rows.append([text(cell.value) for cell in cells])
        finally:
            book.close()
    elif path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            sample = stream.read(8192)
            stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            rows = []
            for index, row in enumerate(csv.reader(stream, dialect)):
                if index > MAX_ROWS:
                    raise ValueError(f"Máximo {MAX_ROWS} casilleros por archivo.")
                rows.append([s[1:] if len(s) > 1 and s[0] == "'" and s[1] in "=+-@\t\r" else s for s in row])
    else:
        raise ValueError("Selecciona un archivo .xlsx o .csv.")
    if not rows or rows[0] != HEADERS:
        raise ValueError("Las columnas no coinciden. Descarga la plantilla desde Inventario.")
    result = []
    for index, row in enumerate(rows[1:], 2):
        if not any(row):
            continue
        if len(row) != len(HEADERS):
            raise ValueError(f"Fila {index}: cantidad de columnas incorrecta.")
        result.append(dict(zip(HEADERS, map(text, row))))
    return result


def prepare_import(rows, snapshot):
    if not rows or len(rows) > MAX_ROWS:
        raise ValueError(f"El archivo debe contener entre 1 y {MAX_ROWS} casilleros.")
    buildings = {b["code"]: b["id"] for b in snapshot["buildings"]}
    blocks = {b["id"]: b for b in snapshot["blocks"]}
    lockers = {locker["number"]: locker for locker in snapshot["lockers"]}
    groups, seen_numbers = {}, set()
    for index, row in enumerate(rows, 2):

        def error(message):
            raise ValueError(f"Fila {index}: {message}")

        def number(key, low, high, integer=False):
            try:
                value = float(row[key])
            except (ValueError, TypeError, KeyError):
                error(f"{key} debe ser numérico.")
            if not math.isfinite(value) or not low <= value <= high or (integer and value != int(value)):
                error(f"{key} debe estar entre {low} y {high}" + (" y ser entero." if integer else "."))
            return int(value) if integer else value

        plan_matches = [p for p in snapshot["plans"] if p["id"] == row["plano"]]
        if not plan_matches:
            plan_matches = [p for p in snapshot["plans"] if p["name"] == row["plano"]]
        if len(plan_matches) != 1:
            error("el plano no existe o el nombre es ambiguo.")
        building_id = buildings.get(row["edificio"])
        if not building_id:
            error("el edificio no existe.")
        if row["lado"] not in ("N", "S", "E", "O") or row["ubicacion"] not in ("EXT", "INT"):
            error("lado o ubicación inválidos.")
        if not row["bloque"] or len(row["bloque"]) > 60:
            error("el nombre del bloque debe tener de 1 a 60 caracteres.")
        p = dict(
            plan_id=plan_matches[0]["id"],
            building_id=building_id,
            label=row["bloque"],
            side=row["lado"],
            location=row["ubicacion"],
            rows=number("filas", 1, 12, True),
            columns=number("columnas", 1, 12, True),
            x=number("x", 0, 1000),
            y=number("y", 0, 1000),
            rotation=number("rotacion", -360, 360),
        )
        existing = (
            blocks.get(row["id_bloque"])
            if row["id_bloque"]
            else next(
                (
                    b
                    for b in blocks.values()
                    if (b["building_id"], b["location"], b["label"]) == (building_id, p["location"], p["label"])
                ),
                None,
            )
        )
        if row["id_bloque"] and not existing:
            error("el ID del bloque no existe; deja el ID vacío solo para crear uno nuevo.")
        version = None
        if existing:
            version = number("version_bloque", 1, 2**31 - 1, True)
            if version != existing["version"]:
                error("el bloque cambió. Exporta nuevamente antes de editar.")
            if any(p[k] != existing[k] for k in ("plan_id", "building_id", "side", "location", "rows", "columns")):
                error("solo se pueden modificar nombre, X, Y y rotación de bloques existentes.")
        key = existing["id"] if existing else (building_id, p["location"], p["label"])
        metadata = (p, version)
        if key not in groups:
            groups[key] = {"metadata": metadata, "existing": existing, "cells": {}}
        group = groups[key]
        if group["metadata"] != metadata:
            error("los datos del bloque deben ser iguales en todas sus filas.")
        cell = (number("fila", 1, p["rows"], True) - 1, number("columna", 1, p["columns"], True) - 1)
        n = row["numero"]
        if not n or len(n) > 40 or n in seen_numbers or cell in group["cells"]:
            error("número vacío, demasiado largo, duplicado o posición repetida.")
        current = lockers.get(n)
        if any(item["number"] == n for item in snapshot.get("archived_lockers", [])):
            error("ese número está eliminado y reservado. Usa Elementos eliminados para restaurarlo.")
        if existing:
            if not current or (current["block_id"], current["row"], current["column"]) != (existing["id"], *cell):
                error(
                    "no se puede cambiar la numeración o posición interna de casilleros existentes mediante importación."
                )
        elif current:
            error("ese número físico ya existe en otro bloque.")
        seen_numbers.add(n)
        group["cells"][cell] = n
    operations, versions = [], {}
    for group in groups.values():
        p, version = group["metadata"]
        existing = group["existing"]
        expected = (
            sum(locker["block_id"] == existing["id"] for locker in lockers.values())
            if existing
            else p["rows"] * p["columns"]
        )
        if len(group["cells"]) != expected:
            raise ValueError(f"Bloque {p['label']}: incluye todos sus casilleros, una fila por posición.")
        existing = group["existing"]
        if existing:
            if any(p[k] != existing[k] for k in ("label", "x", "y", "rotation")):
                operations.append({"action": "block.update", "payload": dict(p, id=existing["id"])})
                versions["block:" + existing["id"]] = version
        else:
            p["numbers"] = [n for _, n in sorted(group["cells"].items())]
            operations.append({"action": "block.create", "payload": p})
    # Catch final name collisions before queuing anything.
    final = {b["id"]: (b["building_id"], b["location"], b["label"]) for b in blocks.values()}
    final.update({b["id"]: (b["building_id"], b["location"], b["label"]) for b in snapshot.get("archived_blocks", [])})
    for index, op in enumerate(operations):
        p = op["payload"]
        final[p.get("id", f"new-{index}")] = (p["building_id"], p["location"], p["label"])
    if len(set(final.values())) != len(final):
        raise ValueError("Dos bloques tendrían el mismo nombre en el mismo edificio y ubicación.")
    if len(operations) > 2000:
        raise ValueError("Máximo 2000 bloques por importación.")
    return operations, versions
